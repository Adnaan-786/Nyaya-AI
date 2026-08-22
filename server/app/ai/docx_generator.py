"""Markdown to DOCX for draftsman output.

A lawyer files a Word document, not Markdown, so the draft has to arrive as something
that opens in Word and can be edited before it goes to court.

Deliberately handles only the subset the draft prompts ask for — `#`/`##`/`###`
headings, `**bold**` spans, and plain paragraphs. Both ends of this pipeline are ours
(the prompt says what to produce, this parses it), so a general Markdown renderer would
be a dependency and a surface area bought for nothing.
"""

import io
import re

from docx import Document

_BOLD_SPAN_RE = re.compile(r"\*\*(.+?)\*\*")

_HEADING_PREFIXES = (("### ", 3), ("## ", 2), ("# ", 1))


def _add_paragraph_with_inline_bold(document: Document, line: str) -> None:
    paragraph = document.add_paragraph()
    position = 0
    for match in _BOLD_SPAN_RE.finditer(line):
        if match.start() > position:
            paragraph.add_run(line[position : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        position = match.end()
    if position < len(line):
        paragraph.add_run(line[position:])


def markdown_to_docx_bytes(markdown_text: str) -> bytes:
    document = Document()

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line:
            continue

        for prefix, level in _HEADING_PREFIXES:
            if line.startswith(prefix):
                document.add_heading(line[len(prefix) :].strip(), level=level)
                break
        else:
            _add_paragraph_with_inline_bold(document, line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
